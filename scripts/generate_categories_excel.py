# -*- coding: utf-8 -*-
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

catalog = [
    {
        "name": "Électroménager",
        "slug": "electromenager",
        "popular": 1,
        "children": [
            {
                "name": "Gros électroménager",
                "slug": "gros-electromenager",
                "children": [
                    "Réfrigérateurs",
                    "Congélateurs",
                    "Cuisinières",
                    "Splits-Climatiseur",
                    "Machine à Laver",
                    "Lave Vaisselle",
                    "Chauffe Eau",
                    "Hotte de Cuisine",
                    "Mini frigo",
                    "Distributeur de glaçons",
                    "Supports Réfrigérateurs",
                    "Accessoires de Cuisinières",
                    "Accessoires de montage splits",
                    "Purificateur d'air",
                    "Ventilateur",
                ],
            },
            {
                "name": "Petits électroménager",
                "slug": "petits-electromenager",
                "children": [
                    "Bouilloire",
                    "Mixeurs",
                    "Air Fryer",
                    "Micro-onde",
                    "Machine à Café",
                    "Capsule machine à café",
                    "Friteuse",
                    "Grille Pain",
                    "Sandwich Maker",
                    "Juice Maker",
                    "Hachoir",
                    "Batteuses",
                    "Four Électrique",
                    "Plaque de Cuisson",
                    "Plaque Chauffante",
                    "Cuiseur de riz",
                    "Thermos",
                    "Barbecue",
                    "Fontaine à Eau",
                ],
            },
            {
                "name": "Appareils ménagers",
                "slug": "appareils-menagers",
                "children": [
                    "Aspirateur",
                    "Fer à Repasser",
                    "Défroisseur",
                    "Tueur d'insecte",
                    "Appareils de cuisine",
                    "Appareils de santé et de beauté",
                    "Luminaires & Éclairage",
                ],
            },
        ],
    },
    {
        "name": "Mobilier de bureau",
        "slug": "bureaux",
        "popular": 1,
        "children": [
            "Accessoires Divers",
            "Banquettes",
            "Bureaux",
            "Chaise de Cérémonie",
            "Chauffeuse de Bureau",
            "Caissons",
            "Chaises de Bureau",
            "Coffres forts",
            "Comptoir de réception",
            "Fauteuils de Bureau",
            "Lampes de bureau",
            "Meubles de rangement",
            "Rangements métalliques",
            "Pupitres de conférence",
            "Salons de bureau",
            "Station de Travail",
            "Table de Réunion",
            "Tables de Bureau",
            "Table Basse de Bureau",
            "Table Coin de Bureau",
        ],
    },
    {
        "name": "Mobilier de maison",
        "slug": "meubles",
        "popular": 1,
        "children": [
            "Décorations de maison",
            "Tables de coin",
            "Salon en cuir",
            "Tables téléviseur",
            "Banquette Chambre à Coucher",
            "Matelas",
            "Salon en tissu",
            "Tables à manger",
            "Buffet",
            "Meubles TV",
            "Tables basses",
            "Meubles chambres à coucher",
            "Chaises table à manger",
            "Consoles",
            "Vaisseliers",
            "Canapé",
            "Chauffeuse 1 Place",
            "Chambres à coucher adulte",
            "Coiffeuse de lit",
            "Lampe",
            "Chambres à coucher enfant",
            "Meubles de jardin",
            "Armoire Chambre à Coucher",
            "Lit",
            "Salon Angle en Tissu",
            "Salon Angle en cuir",
        ],
    },
    {
        "name": "Électronique & TV",
        "slug": "tv-audio",
        "popular": 0,
        "children": [
            "Téléviseur",
            "Barre De Son",
            "Home Théâtre-Cinéma",
            "Support Téléviseur",
            "Casque",
            "Microphones",
            "Tablettes",
            "Téléphones et Accessoires",
            "Radio",
            "Régulateur de Tension",
            "Lampe led",
        ],
    },
    {
        "name": "Destockage",
        "slug": "destockage",
        "popular": 0,
        "children": [],
    },
    {
        "name": "Reconditionné",
        "slug": "reconditionne",
        "popular": 0,
        "children": [],
    },
]

GROUP_SLUGS = {"gros-electromenager", "petits-electromenager", "appareils-menagers"}


def slugify(name: str) -> str:
    s = unicodedata.normalize("NFKD", name)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = s.lower()
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s or "categorie"


def build_rows():
    rows = []
    for i, root in enumerate(catalog):
        rows.append(
            {
                "niveau": 1,
                "name": root["name"],
                "slug": root["slug"],
                "parent_slug": "",
                "display_order": (i + 1) * 10,
                "is_active": 1,
                "is_popular": root["popular"],
                "description": f"{root['name']} — large choix chez DK HOMETECH Dakar.",
            }
        )
        for j, child in enumerate(root["children"]):
            if isinstance(child, dict):
                rows.append(
                    {
                        "niveau": 2,
                        "name": child["name"],
                        "slug": child["slug"],
                        "parent_slug": root["slug"],
                        "display_order": (j + 1) * 10,
                        "is_active": 1,
                        "is_popular": 0,
                        "description": f"{child['name']} — DK HOMETECH Dakar.",
                    }
                )
                for k, leaf in enumerate(child["children"]):
                    leaf_slug = slugify(leaf)
                    if leaf_slug == child["slug"]:
                        leaf_slug = f"{child['slug']}-items"
                    rows.append(
                        {
                            "niveau": 3,
                            "name": leaf,
                            "slug": leaf_slug,
                            "parent_slug": child["slug"],
                            "display_order": (k + 1) * 10,
                            "is_active": 1,
                            "is_popular": 0,
                            "description": "",
                        }
                    )
            else:
                leaf_slug = slugify(child)
                if leaf_slug == root["slug"]:
                    leaf_slug = f"{root['slug']}-items"
                rows.append(
                    {
                        "niveau": 2,
                        "name": child,
                        "slug": leaf_slug,
                        "parent_slug": root["slug"],
                        "display_order": (j + 1) * 10,
                        "is_active": 1,
                        "is_popular": 0,
                        "description": "",
                    }
                )
    return rows


def is_leaf(row: dict) -> bool:
    if row["niveau"] == 3:
        return True
    if row["niveau"] == 1:
        return row["slug"] in {"destockage", "reconditionne"}
    if row["niveau"] == 2:
        return row["slug"] not in GROUP_SLUGS
    return False


def main() -> None:
    rows = build_rows()
    product_rows = []
    for r in rows:
        if not is_leaf(r):
            continue
        product_rows.append(
            {
                "sku": "",
                "name": "",
                "slug": "",
                "category_slug": r["slug"],
                "category_name": r["name"],
                "brand_slug": "",
                "short_description": "",
                "description": "",
                "price": "",
                "promo_price": "",
                "stock_quantity": "",
                "condition": "neuf",
                "is_clearance": 0,
                "is_customizable": 0,
                "status": "draft",
                "meta_title": "",
                "meta_description": "",
                "image_urls": "",
            }
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Categories"
    headers = [
        "niveau",
        "name",
        "slug",
        "parent_slug",
        "display_order",
        "is_active",
        "is_popular",
        "description",
    ]
    header_fill = PatternFill("solid", fgColor="F97316")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )
    level_fills = {
        1: PatternFill("solid", fgColor="FFF7ED"),
        2: PatternFill("solid", fgColor="FFFFFF"),
        3: PatternFill("solid", fgColor="F9FAFB"),
    }

    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for i, r in enumerate(rows, 2):
        for col, h in enumerate(headers, 1):
            cell = ws.cell(i, col, r[h])
            cell.border = thin
            cell.fill = level_fills.get(r["niveau"], level_fills[2])
            if h == "name" and r["niveau"] == 1:
                cell.font = Font(bold=True)

    for i, w in enumerate([10, 36, 36, 28, 14, 12, 12, 50], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:H{len(rows) + 1}"
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Produits_a_remplir")
    p_headers = [
        "sku",
        "name",
        "slug",
        "category_slug",
        "category_name",
        "brand_slug",
        "short_description",
        "description",
        "price",
        "promo_price",
        "stock_quantity",
        "condition",
        "is_clearance",
        "is_customizable",
        "status",
        "meta_title",
        "meta_description",
        "image_urls",
    ]
    for col, h in enumerate(p_headers, 1):
        cell = ws2.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
    for i, r in enumerate(product_rows, 2):
        for col, h in enumerate(p_headers, 1):
            ws2.cell(i, col, r[h]).border = thin
    for i, w in enumerate(
        [12, 32, 32, 32, 32, 16, 28, 40, 12, 12, 12, 12, 12, 14, 12, 28, 36, 50],
        1,
    ):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.auto_filter.ref = f"A1:R{len(product_rows) + 1}"
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Guide")
    guide_lines = [
        "Fichier catégories DK HOMETECH — structure méga-menu",
        "",
        "Contenu",
        "- Feuille Categories : arborescence complète (niveaux 1 → 2 → 3)",
        "- Feuille Produits_a_remplir : une ligne par catégorie feuille, prête pour vos produits + image_urls",
        "",
        "Import catégories en production (recommandé — 1 clic)",
        "1. Admin → Catégories → bouton « Importer suggestion »",
        "2. Cela crée la même arborescence Électroménager / Meubles / Bureaux / TV",
        "3. Ensuite Admin → Navigation pour synchroniser le méga-menu si besoin",
        "",
        "Import produits plus tard",
        "1. Remplir la feuille Produits_a_remplir (name, price, image_urls séparées par | )",
        "2. Enregistrer en CSV (séparateur ;) ou utiliser Admin → Produits → Uploader un CSV",
        "3. La colonne category_slug doit correspondre exactement à la feuille Categories",
        "",
        "Colonnes Categories",
        "niveau = 1 racine, 2 sous-catégorie / groupe, 3 feuille (produits)",
        "parent_slug = slug du parent (vide pour niveau 1)",
        "slug = identifiant unique (ne pas changer après liaison produits)",
        "",
        f"Total catégories : {len(rows)}",
        f"Catégories feuilles (pour produits) : {len(product_rows)}",
    ]
    for i, line in enumerate(guide_lines, 1):
        ws3.cell(i, 1, line)
        if i == 1:
            ws3.cell(i, 1).font = Font(bold=True, size=14, color="F97316")
    ws3.column_dimensions["A"].width = 100

    out = Path(__file__).resolve().parents[1] / "imports" / "categories-dk-hometech.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"OK {out}")
    print(f"categories={len(rows)} leaves={len(product_rows)}")


if __name__ == "__main__":
    main()
